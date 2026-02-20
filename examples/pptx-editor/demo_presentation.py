#!/usr/bin/env python3
"""Generate a demo PowerPoint presentation for testing the PPTX workflow.

Creates work/presentations/demo.pptx with:
- A title slide
- A slide with placeholder text to edit
- A slide with math equation placeholders (to be rendered as LaTeX)
"""
import os

from pptx import Presentation
from pptx.util import Inches

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
REPO_ROOT = os.path.abspath(os.path.join(SCRIPT_DIR, "..", ".."))
OUTPUT_DIR = os.path.join(REPO_ROOT, "work", "presentations")
OUTPUT_PATH = os.path.join(OUTPUT_DIR, "demo.pptx")


def main():
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Slide 1: Title
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "PocketPostdoc Demo Presentation"
    slide.placeholders[1].text = "Edit this presentation using Claude Code"

    # Slide 2: Content with placeholder text
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Project Overview"
    slide.placeholders[1].text = (
        "TODO: Replace this with your project overview.\n\n"
        "- First point\n"
        "- Second point\n"
        "- Third point"
    )

    # Slide 3: Equation placeholders
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Key Equations"
    slide.placeholders[1].text = (
        "The following equations should be rendered in LaTeX:\n\n"
        "Einstein's mass-energy equivalence: E = mc^2\n\n"
        "Euler's identity: e^{i*pi} + 1 = 0\n\n"
        "Ask Claude to render these as LaTeX images!"
    )

    prs.save(OUTPUT_PATH)
    print(f"Demo presentation created: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
